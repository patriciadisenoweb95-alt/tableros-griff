# -*- coding: utf-8 -*-
"""
Lectura de la PLANILLA MAESTRA DE PRESTADORES (formato HUB unificado).

Un archivo Planilla_Prestadores_MM-YYYY.xlsx por mes.
  - Hoja de datos: PRESTADORES
  - Encabezado en la fila 3, datos desde la fila 4
  - El periodo se lee de la celda K1 (formato MM-AAAA); si esta vacia
    se intenta detectar desde el nombre del archivo.

Esta capa convierte el formato nuevo al MISMO DataFrame que esperaba el
resto del pipeline (columnas PRESTADOR, CUIT, Numero de Factura, etc.).
"""

import os
import re
import glob
import pandas as pd
import openpyxl

from .utils import num, norm_nombre, norm_os, norm_cuit, \
                    detectar_periodo, periodo_label, fecha_str, parse_fecha


COLS_OUT = ["PRESTADOR", "CUIT", "Numero de Factura", "FECHA_FACT",
            "OBRA_SOCIAL", "OBRA_SOCIAL_RAW", "A_PAGAR", "PERIODO",
            "PERIODO_CODIGO", "ARCHIVO_ORIGEN"]

COLS_REQ = ["PRESTADOR", "CUIT", "NUMERO_COMPROBANTE", "FECHA_COMPROBANTE",
            "OBRA_SOCIAL", "A_PAGAR", "ESTADO"]


def _leer_periodo(path):
    """Lee el periodo de la celda K1 (MM-AAAA). Fallback: nombre del archivo."""
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb["PRESTADORES"] if "PRESTADORES" in wb.sheetnames else wb.active
        k1 = ws["K1"].value
        wb.close()
        if k1:
            m = re.match(r"^(0?[1-9]|1[0-2])[-_/](20\d{2})$", str(k1).strip())
            if m:
                return f"{int(m.group(1)):02d}-{m.group(2)}"
    except Exception:
        pass
    return detectar_periodo(os.path.basename(path))


def leer_planilla(path):
    """Lee UNA planilla maestra. Devuelve un DataFrame con COLS_OUT."""
    fname = os.path.basename(path)
    periodo = _leer_periodo(path)
    if not periodo:
        raise ValueError(f"No detecto periodo (ni celda K1 ni nombre): {fname}")

    df = pd.read_excel(path, sheet_name="PRESTADORES", header=2, dtype=str,
                       engine="openpyxl")
    df.columns = [str(c).strip() for c in df.columns]

    faltan = [c for c in COLS_REQ if c not in df.columns]
    if faltan:
        raise ValueError(f"Faltan columnas en {fname}: {', '.join(faltan)}")

    s = df["PRESTADOR"].astype(str).str.strip().str.upper()
    invalidos = {"", "NAN", "NONE", "TOTAL", "TOTALES", "NULL"}
    df = df[~s.isin(invalidos)].copy()

    out = pd.DataFrame()
    out["PRESTADOR"] = df["PRESTADOR"].apply(norm_nombre)
    out["CUIT"] = df["CUIT"].apply(norm_cuit)
    out["Numero de Factura"] = df["NUMERO_COMPROBANTE"].astype(str).str.strip()
    out["FECHA_FACT"] = df["FECHA_COMPROBANTE"].apply(
        lambda x: fecha_str(parse_fecha(x)))
    raw_os = df["OBRA_SOCIAL"].astype(str).str.strip()
    out["OBRA_SOCIAL_RAW"] = raw_os
    out["OBRA_SOCIAL"] = raw_os.apply(lambda x: norm_os(x) or "")
    out["A_PAGAR"] = df["A_PAGAR"].apply(num)
    out["PERIODO"] = periodo_label(periodo)
    out["PERIODO_CODIGO"] = periodo
    out["ARCHIVO_ORIGEN"] = fname
    return out[COLS_OUT].reset_index(drop=True)


def leer_carpeta_planillas(carpeta_planillas):
    """Lee todas las Planilla_Prestadores_*.xlsx de la carpeta y las une.
    Devuelve (df_unificado, info_dict)."""
    archivos = glob.glob(os.path.join(carpeta_planillas,
                                      "Planilla_Prestadores_*.xlsx"))
    archivos = [a for a in archivos
                if not os.path.basename(a).lower().startswith("~$")
                and "TEMPLATE" not in os.path.basename(a).upper()]

    dfs, procesados, ignorados = [], [], []
    for path in sorted(archivos):
        try:
            df = leer_planilla(path)
            dfs.append(df)
            procesados.append({
                "archivo": os.path.basename(path),
                "periodo": df["PERIODO"].iloc[0] if len(df) else "",
                "facturas": len(df),
                "monto": df["A_PAGAR"].sum(),
            })
        except Exception as e:
            ignorados.append({"archivo": os.path.basename(path),
                              "motivo": str(e)})

    if dfs:
        df_total = pd.concat(dfs, ignore_index=True)
    else:
        df_total = pd.DataFrame(columns=COLS_OUT)

    return df_total, {"procesados": procesados, "ignorados": ignorados}
