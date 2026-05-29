# -*- coding: utf-8 -*-
"""
Lectura del archivo impuestos.xlsx (una fila por mes con IIBB/IVA/Municipalidad).
Si el archivo no existe, devuelve un dict vacio (impuestos = 0 para todo).
"""

import os
import openpyxl

from .utils import num, detectar_periodo


def leer_impuestos(path):
    """Lee impuestos.xlsx. Devuelve {periodo: {iibb, iva, munic}} e info."""
    resultado = {}
    info = {"procesados": [], "ignorados": []}

    if not os.path.isfile(path):
        info["ignorados"].append({
            "archivo": "impuestos.xlsx",
            "motivo": "No existe el archivo (se asumen impuestos = 0)",
        })
        return resultado, info

    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
    except Exception as e:
        info["ignorados"].append({
            "archivo": "impuestos.xlsx", "motivo": f"No pude abrirlo: {e}"})
        return resultado, info

    # Mapear columnas por header
    header = [str(c.value).strip().upper() if c.value else ""
              for c in ws[1]]
    col_periodo = col_iibb = col_iva = col_munic = None
    for i, h in enumerate(header):
        if "PERIODO" in h:
            col_periodo = i
        elif "IIBB" in h or "INGRESOS BRUTOS" in h:
            col_iibb = i
        elif h == "IVA":
            col_iva = i
        elif "MUNIC" in h:
            col_munic = i

    if col_periodo is None:
        info["ignorados"].append({
            "archivo": "impuestos.xlsx",
            "motivo": "No encontre la columna PERIODO"})
        return resultado, info

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[col_periodo] is None:
            continue
        periodo_raw = str(row[col_periodo]).strip()
        # Aceptar "MM-YYYY" directo o detectarlo
        periodo = periodo_raw if len(periodo_raw) == 7 and periodo_raw[2] == "-" \
            else detectar_periodo(periodo_raw)
        if not periodo:
            continue
        iibb = num(row[col_iibb]) if col_iibb is not None else 0.0
        iva = num(row[col_iva]) if col_iva is not None else 0.0
        munic = num(row[col_munic]) if col_munic is not None else 0.0
        resultado[periodo] = {"iibb": iibb, "iva": iva, "munic": munic}
        info["procesados"].append({
            "periodo": periodo,
            "total": iibb + iva + munic,
        })

    return resultado, info
