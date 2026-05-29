# -*- coding: utf-8 -*-
"""
Lectura de la PLANILLA MAESTRA DE PRESTADORES (formato HUB unificado) para
el Tablero Estado de Resultado. Es la MISMA planilla que usa Cuentas a Pagar.
Hoja PRESTADORES, encabezado en la fila 3, periodo en la celda K1.
"""

import os
import re
import glob
import openpyxl
from datetime import datetime
from collections import defaultdict

from .utils import num, norm_os, norm_categoria, norm_nombre, norm_cuit, \
                    detectar_periodo, periodo_label, fecha_iso

HEADER_ROW = 3


def _leer_periodo(path):
    """Lee el periodo de la celda K1 (MM-AAAA). Fallback: nombre del archivo."""
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb["PRESTADORES"] if "PRESTADORES" in wb.sheetnames else wb.active
        k1 = ws["K1"].value
        wb.close()
        if k1:
            m = re.match(r"^(0?[1-9]|1[0-2])[-_/](20\d{2})$", str(k1).strip())
            if m:
                return "{:02d}-{}".format(int(m.group(1)), m.group(2))
    except Exception:
        pass
    return detectar_periodo(os.path.basename(path))


def leer_planilla(path):
    """Lee UNA planilla maestra. Devuelve dict con periodo + facturas."""
    fname = os.path.basename(path)
    periodo = _leer_periodo(path)
    if not periodo:
        raise ValueError("No detecto periodo (ni K1 ni nombre): " + fname)

    wb = openpyxl.load_workbook(path, data_only=True)
    if "PRESTADORES" not in wb.sheetnames:
        raise ValueError("No tiene hoja PRESTADORES: " + fname)
    ws = wb["PRESTADORES"]

    invoices = []
    for row in ws.iter_rows(min_row=HEADER_ROW + 1, max_row=ws.max_row,
                            values_only=True):
        r = list(row[:14]) + [None] * (14 - len(row[:14]))
        prest, cuit, fact, fecha = r[1], r[2], r[4], r[5]
        periodo_prest, tipo, os_v = r[6], r[7], r[8]
        factura, admin, medico, apagar = r[10], r[11], r[12], r[13]
        if not prest:
            continue
        if str(prest).strip().upper() in ("", "NAN", "NONE", "TOTAL",
                                          "TOTALES", "NULL"):
            continue
        fecha_pago = fecha_iso(fecha)
        periodo_p = fecha_iso(periodo_prest)
        plazo_real = None
        fp = fecha if isinstance(fecha, datetime) else None
        pp = periodo_prest if isinstance(periodo_prest, datetime) else None
        if fp and pp:
            plazo_real = (fp - pp).days
        invoices.append({
            "p": norm_nombre(prest)[:50] if prest else "",
            "f": (str(fact).strip() if fact else "")[:20],
            "fp": fecha_pago,
            "pe": periodo_p,
            "cat": norm_categoria(tipo),
            "os": norm_os(os_v) if os_v else "",
            "plr": plazo_real,
            "ft": num(factura),
            "da": num(admin),
            "dm": num(medico),
            "ap": num(apagar),
        })
    wb.close()
    return {"periodo": periodo, "month": periodo_label(periodo),
            "archivo": fname, "invoices": invoices}


def agregar_planilla(planilla):
    """Toma una planilla leida y calcula todos los agregados del mes."""
    inv = planilla["invoices"]

    # Por prestador
    prest = defaultdict(lambda: {"facturas": 0, "ft": 0.0, "da": 0.0,
                                 "dm": 0.0, "ap": 0.0, "os": set()})
    for i in inv:
        p = prest[i["p"]]
        p["facturas"] += 1
        p["ft"] += i["ft"]; p["da"] += i["da"]
        p["dm"] += i["dm"]; p["ap"] += i["ap"]
        if i["os"]:
            p["os"].add(i["os"])
    prest_list = []
    for name, v in prest.items():
        prest_list.append({
            "p": name, "facturas": v["facturas"],
            "ft": round(v["ft"], 2), "da": round(v["da"], 2),
            "dm": round(v["dm"], 2), "ap": round(v["ap"], 2),
            "os": sorted(v["os"]),
            "deb_pct": round((v["da"] + v["dm"]) / v["ft"] * 100, 2)
                        if v["ft"] else 0,
        })
    prest_list.sort(key=lambda x: -x["ap"])

    # Por categoria
    by_cat = defaultdict(lambda: {"ft": 0.0, "ap": 0.0, "n": 0})
    for i in inv:
        c = i["cat"]
        by_cat[c]["ft"] += i["ft"]; by_cat[c]["ap"] += i["ap"]
        by_cat[c]["n"] += 1
    cat_list = [{"cat": c, "ft": round(v["ft"], 2), "ap": round(v["ap"], 2),
                 "n": v["n"]} for c, v in by_cat.items()]
    cat_list.sort(key=lambda x: -x["ap"])

    # Por OS
    by_os = defaultdict(lambda: {"ft": 0.0, "ap": 0.0, "n": 0})
    for i in inv:
        o = i["os"] or "?"
        by_os[o]["ft"] += i["ft"]; by_os[o]["ap"] += i["ap"]
        by_os[o]["n"] += 1
    os_list = [{"os": o, "ft": round(v["ft"], 2), "ap": round(v["ap"], 2),
                "n": v["n"]} for o, v in by_os.items()]
    os_list.sort(key=lambda x: -x["ap"])

    # Heatmap categoria x OS
    matrix = defaultdict(lambda: defaultdict(float))
    for i in inv:
        matrix[i["cat"]][i["os"] or "?"] += i["ap"]
    heatmap = {c: {o: round(v, 2) for o, v in d.items()}
               for c, d in matrix.items()}

    # Plazos
    plazos = [i["plr"] for i in inv if i["plr"] is not None]
    plazo_stats = None
    if plazos:
        plazo_stats = {
            "n": len(plazos),
            "avg": round(sum(plazos) / len(plazos), 1),
            "min": min(plazos), "max": max(plazos),
            "sin_periodo": len([i for i in inv if i["pe"] is None]),
        }

    totals = {
        "n_facturas": len(inv),
        "n_prestadores": len(prest),
        "ft_total": round(sum(i["ft"] for i in inv), 2),
        "ap_total": round(sum(i["ap"] for i in inv), 2),
        "da_total": round(sum(i["da"] for i in inv), 2),
        "dm_total": round(sum(i["dm"] for i in inv), 2),
        "factura_promedio": round(sum(i["ap"] for i in inv) / len(inv), 2)
                             if inv else 0,
    }

    return {
        "periodo": planilla["periodo"],
        "month": planilla["month"],
        "archivo": planilla["archivo"],
        "totals": totals,
        "categorias": cat_list,
        "obras_sociales": os_list,
        "heatmap": heatmap,
        "prestadores": prest_list,
        "plazos": plazo_stats,
        "invoices": inv,
    }

def leer_carpeta_costos(carpeta):
    """Lee todas las Planilla_Prestadores_*.xlsx de la carpeta.
    Devuelve (dict_por_periodo, info)."""
    archivos = glob.glob(os.path.join(carpeta, "Planilla_Prestadores_*.xlsx"))
    archivos = [a for a in archivos
                if not os.path.basename(a).startswith("~$")
                and "TEMPLATE" not in os.path.basename(a).upper()]
    por_periodo = {}
    procesados = []
    ignorados = []
    for path in sorted(archivos):
        try:
            planilla = leer_planilla(path)
            agg = agregar_planilla(planilla)
            por_periodo[agg["periodo"]] = agg
            procesados.append({
                "archivo": os.path.basename(path),
                "periodo": agg["month"],
                "facturas": agg["totals"]["n_facturas"],
                "a_pagar": agg["totals"]["ap_total"],
            })
        except Exception as e:
            ignorados.append({
                "archivo": os.path.basename(path),
                "motivo": str(e),
            })
    info = {"procesados": procesados, "ignorados": ignorados}
    return por_periodo, info
