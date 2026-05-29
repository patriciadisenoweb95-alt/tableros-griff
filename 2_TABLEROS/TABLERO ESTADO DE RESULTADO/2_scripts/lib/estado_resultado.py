# -*- coding: utf-8 -*-
"""
Lectura de los Excel mensuales del Estado de Resultado.
Cada archivo = un mes. Detecta el periodo desde el nombre del archivo.

De cada Excel extrae, por obra social:
  - cantidad de afiliados, costo de capita, subtotal capita
  - refuerzo capita, notas de credito/debito, ingresos extras
  - ingresos netos
  - total egresos directos (costos prestacionales)
  - costos fijos (recalculados por participacion de ingresos + MUSCARELO directo)
"""

import os
import glob
import openpyxl

from .config import ER_COLS, ER_CONCEPTOS, ER_LINEAS_FIJAS, ER_LINEA_MUSCARELO, \
                    OS_LIST, OSPIF_SUBS
from .utils import num, detectar_periodo, periodo_label


def _leer_excel_crudo(path):
    """Lee la hoja Estado_Resultado y devuelve la lista de filas (tuplas)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    # Buscar la hoja: preferir "Estado_Resultado", sino la primera
    sheet = None
    for sn in wb.sheetnames:
        if "estado" in sn.lower() and "resultado" in sn.lower():
            sheet = sn
            break
    if sheet is None:
        sheet = wb.sheetnames[0]
    ws = wb[sheet]
    return list(ws.iter_rows(values_only=True))


def _valor_celda(row, col_idx):
    """Devuelve el valor numerico de una celda de la fila, 0 si vacia."""
    if col_idx >= len(row):
        return 0.0
    return num(row[col_idx])


def leer_estado_resultado(path):
    """Lee UN Excel de Estado de Resultado. Devuelve un dict con los datos
    del mes ya normalizados y los costos fijos recalculados por OS."""
    fname = os.path.basename(path)
    periodo = detectar_periodo(fname)
    if not periodo:
        raise ValueError(f"No detecto periodo en el nombre: {fname}")

    rows = _leer_excel_crudo(path)

    # 1) Extraer cada concepto (primera coincidencia top-down)
    conceptos = {k: {} for k in ER_CONCEPTOS}
    todas_cols = dict(ER_COLS)  # incluye subsegmentos OSPIF + TOTAL

    for row in rows:
        if not row or row[0] is None:
            continue
        texto = str(row[0]).strip()
        if "Margen" in texto:
            continue
        for clave, buscar in ER_CONCEPTOS.items():
            if buscar in texto and not conceptos[clave]:
                for os_name, col in todas_cols.items():
                    conceptos[clave][os_name] = _valor_celda(row, col)
                break

    # 2) Lineas de costos fijos: sumar el TOTAL de cada linea (todo el mes)
    fijos_total_lineas = 0.0
    muscarelo_total = 0.0
    for row in rows:
        if not row or row[0] is None:
            continue
        texto = str(row[0]).strip()
        if texto.upper().startswith("TOTAL"):
            continue
        for buscar in ER_LINEAS_FIJAS:
            if buscar in texto:
                fijos_total_lineas += _valor_celda(row, ER_COLS["TOTAL"])
                break
        if ER_LINEA_MUSCARELO in texto.upper():
            muscarelo_total = _valor_celda(row, ER_COLS["TOTAL"])

    # 3) Armar resumen por obra social
    total_ingresos = conceptos["ingresos_netos"].get("TOTAL", 0.0)

    obras = []
    for os_name in OS_LIST:
        ingresos = conceptos["ingresos_netos"].get(os_name, 0.0)
        share = ingresos / total_ingresos if total_ingresos else 0.0
        # Costos fijos: parte ponderada + MUSCARELO si es OSSURRBAC
        fijos = fijos_total_lineas * share
        if os_name == "OSSURRBAC":
            fijos += muscarelo_total
        directos = conceptos["total_directos"].get(os_name, 0.0)
        obras.append({
            "code": os_name,
            "afiliados": int(conceptos["afiliados"].get(os_name, 0)),
            "capita_unit": conceptos["costo_capita"].get(os_name, 0.0),
            "subtotal_capita": conceptos["subtotal_capita"].get(os_name, 0.0),
            "refuerzo_capita": conceptos["refuerzo_capita"].get(os_name, 0.0),
            "notas_credito": conceptos["notas_credito"].get(os_name, 0.0),
            "ingresos_extras": conceptos["ingresos_extras"].get(os_name, 0.0),
            "ingresos_netos": ingresos,
            "income_share": share,
            "costos_prestacionales": directos,
            "costos_fijos": fijos,
        })

    # 4) Sub-segmentos OSPIF (solo para tabla de detalle)
    ospif_breakdown = []
    for sub in OSPIF_SUBS:
        afil = int(conceptos["afiliados"].get(sub, 0))
        ing = conceptos["ingresos_netos"].get(sub, 0.0)
        if afil == 0 and ing == 0:
            continue
        ospif_breakdown.append({
            "code": sub,
            "label": sub.replace("OSPIF_", ""),
            "afiliados": afil,
            "capita_unit": conceptos["costo_capita"].get(sub, 0.0),
            "ingresos_netos": ing,
            "subtotal_capita": conceptos["subtotal_capita"].get(sub, 0.0),
        })

    # 5) Promedio ponderado de capita (solo OS con capita > 0)
    con_capita = [(o["capita_unit"], o["afiliados"]) for o in obras
                  if o["capita_unit"] > 0 and o["afiliados"] > 0]
    if con_capita:
        suma_pond = sum(c * a for c, a in con_capita)
        suma_afil = sum(a for _, a in con_capita)
        capita_prom = suma_pond / suma_afil if suma_afil else 0.0
    else:
        capita_prom = 0.0

    # 6) Totales del mes (costos fijos recalculados = lineas + MUSCARELO)
    fijos_total = fijos_total_lineas + muscarelo_total
    totals = {
        "afiliados": int(conceptos["afiliados"].get("TOTAL", 0)),
        "ingresos_netos": total_ingresos,
        "ingresos_extras": conceptos["ingresos_extras"].get("TOTAL", 0.0),
        "subtotal_capita": conceptos["subtotal_capita"].get("TOTAL", 0.0),
        "refuerzo_capita": conceptos["refuerzo_capita"].get("TOTAL", 0.0),
        "notas_credito": conceptos["notas_credito"].get("TOTAL", 0.0),
        "costos_prestacionales": conceptos["total_directos"].get("TOTAL", 0.0),
        "costos_fijos": fijos_total,
        "capita_promedio_ponderado": capita_prom,
    }

    return {
        "periodo": periodo,
        "month": periodo_label(periodo),
        "archivo": fname,
        "totals": totals,
        "obras_sociales": obras,
        "ospif_breakdown": ospif_breakdown,
    }


def leer_carpeta_estado_resultado(carpeta):
    """Lee todos los Excel de la carpeta. Devuelve (dict_por_periodo, info)."""
    archivos = []
    for ext in ("*.xlsx", "*.xls", "*.xlsm"):
        archivos.extend(glob.glob(os.path.join(carpeta, ext)))
    archivos = [a for a in archivos
                if not os.path.basename(a).startswith("~$")
                and not os.path.basename(a).startswith(".")]

    por_periodo = {}
    procesados = []
    ignorados = []
    for path in sorted(archivos):
        try:
            data = leer_estado_resultado(path)
            por_periodo[data["periodo"]] = data
            procesados.append({
                "archivo": os.path.basename(path),
                "periodo": data["month"],
                "ingresos": data["totals"]["ingresos_netos"],
                "afiliados": data["totals"]["afiliados"],
            })
        except Exception as e:
            ignorados.append({
                "archivo": os.path.basename(path),
                "motivo": str(e),
            })

    info = {"procesados": procesados, "ignorados": ignorados}
    return por_periodo, info
