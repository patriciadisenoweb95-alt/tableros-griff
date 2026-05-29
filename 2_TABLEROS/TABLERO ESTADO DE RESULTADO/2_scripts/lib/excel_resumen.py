# -*- coding: utf-8 -*-
"""
Genera TABLERO_Resumen.xlsx con todo consolidado en hojas:
  - Estado_Resultado : una fila por OS x mes con el estado de resultado completo
  - Costos_x_Categoria : costos prestacionales por categoria x mes
  - Top_Prestadores : ranking de prestadores por mes
  - Facturas : detalle completo de todas las facturas de todos los meses
  - Impuestos : los impuestos cargados por mes
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

AZUL = "0A1F8C"
CYAN = "0BC8FA"
GRIS = "F4F6FB"

_HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
_HEAD_FILL = PatternFill("solid", fgColor=AZUL)
_MONEY = '#,##0.00'
_PCT = '0.0%'


def _escribir_hoja(ws, headers, rows, money_cols=None, pct_cols=None):
    """Escribe una hoja con header estilizado y filas. money_cols/pct_cols
    son listas de indices (1-based) con formato numerico."""
    money_cols = money_cols or []
    pct_cols = pct_cols or []
    ws.append(headers)
    for cell in ws[1]:
        cell.font = _HEAD_FONT
        cell.fill = _HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append(r)
    # Formato de columnas numericas
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            if cell.column in money_cols:
                cell.number_format = _MONEY
            elif cell.column in pct_cols:
                cell.number_format = _PCT
    # Ancho de columnas
    for i, h in enumerate(headers, 1):
        col = openpyxl.utils.get_column_letter(i)
        maxlen = len(str(h))
        for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
            v = row[0].value
            if v is not None:
                maxlen = max(maxlen, len(str(v)))
        ws.column_dimensions[col].width = min(maxlen + 3, 45)
    ws.freeze_panes = "A2"


def generar_excel(payload, output_path):
    """Escribe TABLERO_Resumen.xlsx."""
    ER = payload["estado_resultado"]
    CP = payload["costos_prestacionales"]
    periodos = payload["periodos"]

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # --- Hoja 1: Estado de Resultado por OS x mes ---
    ws = wb.create_sheet("Estado_Resultado")
    headers = ["Periodo", "Obra Social", "Afiliados", "Capita Unit.",
               "Subtotal Capita", "Refuerzo", "Notas C/D", "Ingresos Extras",
               "Ingresos Netos", "% Ingresos", "Costos Prestacionales",
               "Costos Fijos", "Resultado Antes Imp.", "Impuestos",
               "Resultado Neto", "Margen"]
    rows = []
    for p in periodos:
        m = ER.get(p)
        if not m:
            continue
        for o in m["obras_sociales"]:
            rows.append([
                m["month"], o["code"], o["afiliados"], o["capita_unit"],
                o["subtotal_capita"], o["refuerzo_capita"], o["notas_credito"],
                o["ingresos_extras"], o["ingresos_netos"], o["income_share"],
                o["costos_prestacionales"], o["costos_fijos"],
                o["resultado_antes"], o["impuestos"], o["resultado_neto"],
                o["margen"],
            ])
        # Fila TOTAL del mes
        t = m["totals"]
        rows.append([
            m["month"], "TOTAL", t["afiliados"], "", t["subtotal_capita"],
            t["refuerzo_capita"], t["notas_credito"], t["ingresos_extras"],
            t["ingresos_netos"], 1.0, t["costos_prestacionales"],
            t["costos_fijos"], t["resultado_antes"], t["impuestos"],
            t["resultado_neto"], t["margen"],
        ])
    _escribir_hoja(ws, headers, rows,
                   money_cols=[4, 5, 6, 7, 8, 9, 11, 12, 13, 14, 15],
                   pct_cols=[10, 16])

    # --- Hoja 2: Costos por Categoria x mes ---
    ws = wb.create_sheet("Costos_x_Categoria")
    headers = ["Periodo", "Categoria", "# Facturas", "Facturado", "A Pagar"]
    rows = []
    for p in periodos:
        m = CP.get(p)
        if not m:
            continue
        for c in m["categorias"]:
            rows.append([m["month"], c["cat"], c["n"], c["ft"], c["ap"]])
    _escribir_hoja(ws, headers, rows, money_cols=[4, 5])

    # --- Hoja 3: Top Prestadores x mes ---
    ws = wb.create_sheet("Top_Prestadores")
    headers = ["Periodo", "Prestador", "# Facturas", "Facturado",
               "Debito Adm.", "Debito Med.", "Debitos %", "A Pagar", "OS"]
    rows = []
    for p in periodos:
        m = CP.get(p)
        if not m:
            continue
        for pr in m["prestadores"]:
            rows.append([
                m["month"], pr["p"], pr["facturas"], pr["ft"], pr["da"],
                pr["dm"], pr["deb_pct"] / 100.0, pr["ap"], ", ".join(pr["os"]),
            ])
    _escribir_hoja(ws, headers, rows, money_cols=[4, 5, 6, 8], pct_cols=[7])

    # --- Hoja 4: Facturas (detalle completo) ---
    ws = wb.create_sheet("Facturas")
    headers = ["Periodo Pago", "Prestador", "Factura", "Fecha Pago",
               "Periodo Prestacion", "Categoria", "Obra Social",
               "Plazo (dias)", "Facturado", "Debito Adm.", "Debito Med.",
               "A Pagar"]
    rows = []
    for p in periodos:
        m = CP.get(p)
        if not m:
            continue
        for i in m["invoices"]:
            rows.append([
                m["month"], i["p"], i["f"], i["fp"], i["pe"], i["cat"],
                i["os"], i["plr"], i["ft"], i["da"], i["dm"], i["ap"],
            ])
    _escribir_hoja(ws, headers, rows, money_cols=[9, 10, 11, 12])

    # --- Hoja 5: Impuestos ---
    ws = wb.create_sheet("Impuestos")
    headers = ["Periodo", "IIBB", "IVA", "Municipalidad", "Total"]
    rows = []
    for p in periodos:
        m = ER.get(p)
        if not m:
            continue
        imp = m.get("impuestos_detalle", {"iibb": 0, "iva": 0, "munic": 0})
        tot = imp["iibb"] + imp["iva"] + imp["munic"]
        rows.append([m["month"], imp["iibb"], imp["iva"], imp["munic"], tot])
    _escribir_hoja(ws, headers, rows, money_cols=[2, 3, 4, 5])

    wb.save(output_path)
    return output_path
