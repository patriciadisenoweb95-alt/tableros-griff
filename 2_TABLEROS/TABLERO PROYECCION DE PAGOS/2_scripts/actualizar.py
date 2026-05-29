# -*- coding: utf-8 -*-
"""
actualizar.py - TABLERO DE PROYECCION DE PAGOS - Griff Salud

Dos pestanas:
  PROYECCION  -> simulador de pagos. Elegis que planillas proyectar, una
                 fecha base, y podes ajustar plazo / fecha / metodo de cada
                 factura para jugar con escenarios. La linea de tiempo separa
                 cheque, transferencia y los cheques ya emitidos del banco.
  HISTORICO   -> reporte por prestador de los cheques ya emitidos, cruzado
                 con los meses de prestacion (las planillas).

Salidas (en 3_RESULTADOS/proyeccion/):
  - proyeccion.html         dashboard
  - Proyeccion_Pagos.xlsx   detalle descargable
"""

import os
import re
import sys
import csv
import glob
import json
from datetime import datetime, date, timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(SCRIPT_DIR)
HUB = os.path.dirname(os.path.dirname(ROOT))
DIR_PRESTADORES = os.path.join(HUB, "1_DATOS", "prestadores")
DIR_CHEQUES = os.path.join(HUB, "1_DATOS", "cheques")
DIR_OUT = os.path.join(HUB, "3_RESULTADOS", "proyeccion")
os.makedirs(DIR_OUT, exist_ok=True)

ESTADOS_PROYECCION = {"AUDITADA", "PROYECTADA"}
CHEQUE_CERRADO = ("pagado", "anulado", "rechazado", "repudiado")
HOY = datetime.now().date()


def to_date(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def to_num(v):
    if v is None:
        return 0.0
    s = str(v).strip().replace("$", "").replace(" ", "")
    if not s or s.lower() in ("nan", "none"):
        return 0.0
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def fmt_money(x):
    return "$ " + format(x, ",.2f").replace(",", "@").replace(".", ",").replace("@", ".")


def planilla_label(path):
    """Planilla_Prestadores_03-2026.xlsx -> 03-2026"""
    m = re.search(r"(\d{2})-(\d{4})", os.path.basename(path))
    if m:
        return m.group(1) + "-" + m.group(2)
    return os.path.splitext(os.path.basename(path))[0]


def orden_planilla(lbl):
    m = re.match(r"(\d{2})-(\d{4})", lbl)
    if m:
        return (int(m.group(2)), int(m.group(1)))
    return (9999, 99)


def parse_periodo(motivo):
    """Extrae 'MM/YYYY' del motivo del cheque si dice PRESTACIONES MM/YYYY."""
    m = re.search(r"PRESTACION(?:ES)?\s*(\d{1,2})\s*/\s*(\d{4})",
                  (motivo or "").upper())
    if m:
        return "%02d/%s" % (int(m.group(1)), m.group(2))
    return ""


def leer_facturas():
    """Facturas en estado AUDITADA/PROYECTADA de todas las planillas."""
    facturas = []
    archivos = sorted(glob.glob(os.path.join(DIR_PRESTADORES,
                                             "Planilla_Prestadores_*.xlsx")))
    for path in archivos:
        base = os.path.basename(path)
        if "TEMPLATE" in base.upper() or base.startswith("~$"):
            continue
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
        except Exception:
            continue
        if "PRESTADORES" not in wb.sheetnames:
            wb.close()
            continue
        lbl = planilla_label(path)
        ws = wb["PRESTADORES"]
        for row in ws.iter_rows(min_row=4, values_only=True):
            r = list(row[:25]) + [None] * (25 - len(row[:25]))
            prest = r[1]
            if not prest:
                continue
            if str(prest).strip().upper() in ("", "NAN", "NONE", "TOTAL",
                                              "TOTALES", "NULL"):
                continue
            estado = str(r[14] or "").strip().upper()
            if estado not in ESTADOS_PROYECCION:
                continue
            plazo = None
            if r[9] not in (None, ""):
                try:
                    plazo = int(float(str(r[9]).strip()))
                except (ValueError, TypeError):
                    plazo = None
            ffact = to_date(r[5])
            fecha = to_date(r[23]) or to_date(r[22])
            origen_fecha = "proyectada" if to_date(r[23]) else "sugerida"
            if not fecha:
                if ffact and plazo is not None:
                    fecha = ffact + timedelta(days=plazo)
                    origen_fecha = "calculada"
            if not fecha:
                origen_fecha = "sin fecha"
            monto = to_num(r[24]) or to_num(r[13])
            facturas.append({
                "planilla": lbl,
                "prestador": str(prest).strip(),
                "cuit": str(r[2] or "").strip(),
                "factura": str(r[4] or "").strip(),
                "os": str(r[8] or "").strip() or "SIN OS",
                "estado": estado,
                "metodo": str(r[16] or "").strip() or "(sin metodo)",
                "plazo": plazo,
                "ffact": ffact,
                "fecha": fecha,
                "origen_fecha": origen_fecha,
                "monto": monto,
            })
        wb.close()
    return facturas


def leer_cheques():
    """TODOS los cheques del archivo de Galicia, con detalle completo."""
    cheques = []
    for path in glob.glob(os.path.join(DIR_CHEQUES, "*.csv")):
        data = None
        for enc in ("latin-1", "utf-8-sig", "utf-8"):
            try:
                with open(path, encoding=enc, newline="") as f:
                    data = list(csv.reader(f, delimiter=";"))
                break
            except Exception:
                continue
        if not data or len(data) < 3:
            continue
        for row in data[2:]:
            if len(row) < 8:
                continue
            r = list(row) + [""] * (13 - len(row))
            estado = str(r[7] or "").strip()
            importe = to_num(r[6])
            if importe <= 0:
                continue
            motivo = str(r[12] or "").strip()
            cheques.append({
                "nro": str(r[0] or "").strip(),
                "emitido_a": str(r[1] or "").strip(),
                "cuit": str(r[3] or "").strip(),
                "fecha_pago": to_date(r[4]),
                "fecha_emision": to_date(r[5]),
                "importe": importe,
                "estado": estado,
                "motivo": motivo,
                "periodo": parse_periodo(motivo),
                "pendiente": not any(c in estado.lower()
                                     for c in CHEQUE_CERRADO),
            })
    return cheques


def build_historico(facturas, cheques):
    """Agrupa los cheques por prestador y los cruza con sus meses de
    prestacion (las planillas donde el CUIT tiene facturas)."""
    pm = {}
    for f in facturas:
        cu = f["cuit"]
        if not cu:
            continue
        d = pm.setdefault(cu, {"nombre": f["prestador"],
                               "meses": set(), "facturado": 0.0})
        d["meses"].add(f["planilla"])
        d["facturado"] += f["monto"]
        if not d["nombre"]:
            d["nombre"] = f["prestador"]

    grupos = {}
    for c in cheques:
        cu = c["cuit"] or "(sin cuit)"
        grupos.setdefault(cu, []).append(c)

    out = []
    for cu, chs in grupos.items():
        p = pm.get(cu)
        nombre = (p["nombre"] if p else "") or chs[0]["emitido_a"] \
            or "(sin nombre)"
        meses = sorted(p["meses"], key=orden_planilla) if p else []
        chs_ord = sorted(chs, key=lambda x: x["fecha_pago"] or date(1900, 1, 1))
        out.append({
            "cuit": cu,
            "nombre": nombre,
            "enPlanilla": bool(p),
            "meses": meses,
            "facturado": round(p["facturado"], 2) if p else 0.0,
            "cheques": [{
                "nro": c["nro"],
                "femis": c["fecha_emision"].isoformat() if c["fecha_emision"] else None,
                "fpago": c["fecha_pago"].isoformat() if c["fecha_pago"] else None,
                "m": round(c["importe"], 2),
                "estado": c["estado"],
                "periodo": c["periodo"],
                "motivo": c["motivo"],
            } for c in chs_ord],
        })
    out.sort(key=lambda x: -sum(c["m"] for c in x["cheques"]))
    return out


def main():
    print("=" * 60)
    print(" TABLERO DE PROYECCION DE PAGOS - Griff Salud")
    print("=" * 60)
    facturas = leer_facturas()
    cheques = leer_cheques()

    planillas = sorted({f["planilla"] for f in facturas}, key=orden_planilla)
    pend = [c for c in cheques if c["pendiente"] and c["fecha_pago"]]
    historico = build_historico(facturas, cheques)

    print("  Planillas detectadas: " + (", ".join(planillas) or "(ninguna)"))
    print("  Facturas a proyectar (AUDITADA/PROYECTADA): " + str(len(facturas)))
    print("  Cheques totales: " + str(len(cheques))
          + "  (pendientes: " + str(len(pend)) + ")")
    print("  Prestadores con cheques (historico): " + str(len(historico)))

    total_fact = sum(f["monto"] for f in facturas)
    total_chq = sum(c["importe"] for c in pend)
    print("  Total facturas proyectadas: " + fmt_money(total_fact))
    print("  Total cheques pendientes:   " + fmt_money(total_chq))

    generar_html(facturas, pend, planillas, historico)
    generar_excel(facturas)

    print("  HTML:  " + os.path.join(DIR_OUT, "proyeccion.html"))
    print("  Excel: " + os.path.join(DIR_OUT, "Proyeccion_Pagos.xlsx"))
    print("=" * 60)
    print(" LISTO")


def generar_html(facturas, pend, planillas, historico):
    AZUL = "#1A2D9C"
    CELESTE = "#29ABE2"

    fac_js = [{
        "i": f["planilla"] + "|" + f["cuit"] + "|" + f["factura"],
        "p": f["planilla"],
        "prest": f["prestador"],
        "os": f["os"],
        "plazo": f["plazo"],
        "m": round(f["monto"], 2),
    } for f in facturas]
    chq_js = [{"f": c["fecha_pago"].isoformat(), "m": round(c["importe"], 2)}
              for c in pend]

    resumen = {}
    for f in facturas:
        d = resumen.setdefault(f["planilla"], {"n": 0, "tot": 0.0})
        d["n"] += 1
        d["tot"] += f["monto"]

    checks = []
    for p in planillas:
        d = resumen.get(p, {"n": 0, "tot": 0.0})
        checks.append(
            '<label class="chk"><input type="checkbox" class="plsel" '
            'value="' + p + '" checked> <span class="chk-nom">' + p + '</span>'
            '<span class="chk-meta">' + str(d["n"]) + ' fact &middot; '
            + fmt_money(d["tot"]) + '</span></label>'
        )
    checks_html = "\n".join(checks) or \
        '<span class="muted">No se encontraron planillas.</span>'

    tot_chq = sum(c["importe"] for c in pend)
    chq_chk = (
        '<label class="chk chk-chq"><input type="checkbox" id="chqsel" checked> '
        '<span class="chk-nom">Cheques ya emitidos</span>'
        '<span class="chk-meta">' + str(len(pend)) + ' cheques &middot; '
        + fmt_money(tot_chq) + '</span></label>'
    )

    plantilla = os.path.join(SCRIPT_DIR, "plantilla.html")
    with open(plantilla, encoding="utf-8") as fh:
        html = fh.read()
    html = (html.replace("__AZUL__", AZUL).replace("__CELESTE__", CELESTE)
            .replace("__FECHA__", datetime.now().strftime("%d/%m/%Y %H:%M"))
            .replace("__CHECKS__", checks_html)
            .replace("__CHQ_CHK__", chq_chk)
            .replace("__FAC_JS__", json.dumps(fac_js, ensure_ascii=False))
            .replace("__CHQ_JS__", json.dumps(chq_js, ensure_ascii=False))
            .replace("__PREST_JS__", json.dumps(historico, ensure_ascii=False))
            .replace("__HOY__", HOY.isoformat()))
    with open(os.path.join(DIR_OUT, "proyeccion.html"), "w",
              encoding="utf-8") as f:
        f.write(html)


def generar_excel(facturas):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Proyeccion"
    hdr = ["PLANILLA", "FECHA_PAGO", "PRESTADOR", "CUIT", "NUMERO_FACTURA",
           "OBRA_SOCIAL", "PLAZO", "ESTADO", "METODO_PAGO", "MONTO",
           "ORIGEN_FECHA"]
    azul = PatternFill("solid", start_color="1A2D9C")
    white = Font(name="Arial", bold=True, color="FFFFFF", size=9)
    for i, h in enumerate(hdr, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = azul
        c.font = white
        c.alignment = Alignment(horizontal="center")
    anchos = [11, 13, 32, 13, 18, 14, 8, 12, 14, 16, 13]
    for i, a in enumerate(anchos, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = a
    filas = sorted(facturas, key=lambda x: (x["planilla"],
                                            x["fecha"] or date(2099, 1, 1),
                                            -x["monto"]))
    r = 2
    for f in filas:
        ws.cell(row=r, column=1, value=f["planilla"])
        ws.cell(row=r, column=2,
                value=f["fecha"].strftime("%d/%m/%Y") if f["fecha"] else "")
        ws.cell(row=r, column=3, value=f["prestador"])
        ws.cell(row=r, column=4, value=f["cuit"])
        ws.cell(row=r, column=5, value=f["factura"])
        ws.cell(row=r, column=6, value=f["os"])
        ws.cell(row=r, column=7, value=f["plazo"])
        ws.cell(row=r, column=8, value=f["estado"])
        ws.cell(row=r, column=9, value=f["metodo"])
        cm = ws.cell(row=r, column=10, value=round(f["monto"], 2))
        cm.number_format = '#,##0.00'
        ws.cell(row=r, column=11, value=f["origen_fecha"])
        for col in range(1, 12):
            ws.cell(row=r, column=col).font = Font(name="Arial", size=9)
        r += 1
    ws.cell(row=r, column=9, value="TOTAL").font = Font(name="Arial", bold=True)
    ct = ws.cell(row=r, column=10, value="=SUM(J2:J" + str(r - 1) + ")")
    ct.font = Font(name="Arial", bold=True)
    ct.number_format = '#,##0.00'
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:K" + str(r - 1)
    wb.save(os.path.join(DIR_OUT, "Proyeccion_Pagos.xlsx"))


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        import traceback
        traceback.print_exc()
        sys.exit(1)
